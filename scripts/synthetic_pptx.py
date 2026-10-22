"""Build synthetic PPTX bytes for V4.3 tests and the offline demo.

This module never reads a real course file and never writes to the private store.
"""

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches


def build_synthetic_bytes(
    *, marker: str = "V43_SYNTHETIC_BODY", extra_slide: bool = False
) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "V4.3 Synthetic Deck"
    body = slide.placeholders[1].text_frame
    body.text = marker
    second = body.add_paragraph()
    second.text = "Second synthetic paragraph"
    second.level = 1

    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(4), Inches(4), Inches(1.2)
    ).table
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"

    if extra_slide:
        presentation.slides.add_slide(presentation.slide_layouts[6])

    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
