import hashlib
import json
import struct
import subprocess
import sys
import unittest
import zlib
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

import jsonschema
from pptx import Presentation
from pptx.util import Inches

from scripts.pptx_content import (
    SCHEMA_VERSION,
    PptxContentError,
    dumps_content,
    parse_pptx,
    render_html,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
SCHEMA_PATH = REPO_ROOT / "schemas" / "pptx-content-v1.schema.json"
SCRIPT_PATH = REPO_ROOT / "scripts" / "pptx_content.py"


def _png_1x1() -> bytes:
    """Build a minimal valid 1x1 RGBA PNG without external image libraries."""

    def chunk(kind: bytes, data: bytes) -> bytes:
        return (
            struct.pack(">I", len(data))
            + kind
            + data
            + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
        )

    raw = b"\x00\xff\x00\x00\xff"
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def make_synthetic_pptx(path: Path) -> Path:
    presentation = Presentation()

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Synthetic Title"
    body = slide.placeholders[1].text_frame
    body.text = "First paragraph"
    second = body.add_paragraph()
    second.text = "Second paragraph"
    second.level = 1
    third = body.add_paragraph()
    third.text = "Third paragraph"
    third.level = 2

    table_shape = slide.shapes.add_table(
        2, 2, Inches(1), Inches(4.2), Inches(4), Inches(1.2)
    )
    table = table_shape.table
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    table.cell(1, 0).text = "A2"
    table.cell(1, 1).text = "B2"

    textbox = slide.shapes.add_textbox(Inches(1), Inches(5.6), Inches(4), Inches(0.8))
    textbox.text_frame.text = "Text box note"

    slide.shapes.add_picture(
        BytesIO(_png_1x1()), Inches(6), Inches(4.2), width=Inches(0.5)
    )

    # Blank slide: no shapes, no blocks.
    presentation.slides.add_slide(presentation.slide_layouts[6])

    special = presentation.slides.add_slide(presentation.slide_layouts[1])
    special.shapes.title.text = 'Special <script>alert("x")</script> & "quotes"'
    special_body = special.placeholders[1].text_frame
    special_body.text = 'Text with <tag> & "quotes"'
    special_table = special.shapes.add_table(
        1, 1, Inches(1), Inches(3.5), Inches(3), Inches(1)
    ).table
    special_table.cell(0, 0).text = '<td> & "cell"'

    presentation.save(str(path))
    return path


def _block(slide: dict, block_type: str, role: str | None = None) -> dict:
    for block in slide["blocks"]:
        if block["type"] == block_type and (role is None or block.get("role") == role):
            return block
    raise AssertionError(f"block not found: type={block_type} role={role}")


class PptxContentTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls._tmp = TemporaryDirectory()
        cls.pptx_path = make_synthetic_pptx(Path(cls._tmp.name) / "synthetic-deck.pptx")
        cls.content = parse_pptx(cls.pptx_path)
        cls.json_text = dumps_content(cls.content)
        cls.html_text = render_html(cls.content)
        cls.source_sha = hashlib.sha256(cls.pptx_path.read_bytes()).hexdigest()
        cls.slides = {slide["slide_number"]: slide for slide in cls.content["slides"]}

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_schema_version_sha_and_traceability(self):
        self.assertEqual(self.content["schema_version"], SCHEMA_VERSION)
        self.assertEqual(self.content["source_sha256"], self.source_sha)
        self.assertEqual(len(self.content["slides"]), 3)
        self.assertEqual([slide["slide_number"] for slide in self.content["slides"]], [1, 2, 3])
        self.assertEqual([slide["slide_index"] for slide in self.content["slides"]], [0, 1, 2])

        first = self.slides[1]
        self.assertEqual(first["title"], "Synthetic Title")
        title_block = _block(first, "text", "title")
        self.assertEqual(title_block["shape_index"], 0)
        body_block = _block(first, "text", "body")
        self.assertEqual(body_block["shape_index"], 1)
        self.assertEqual(
            [paragraph["text"] for paragraph in body_block["paragraphs"]],
            ["First paragraph", "Second paragraph", "Third paragraph"],
        )
        self.assertEqual(
            [paragraph["level"] for paragraph in body_block["paragraphs"]], [0, 1, 2]
        )
        self.assertEqual(
            [paragraph["index"] for paragraph in body_block["paragraphs"]], [0, 1, 2]
        )

    def test_table_and_textbox_extraction(self):
        first = self.slides[1]
        table_block = _block(first, "table")
        self.assertEqual(table_block["rows"], [["A1", "B1"], ["A2", "B2"]])

        textbox_block = next(
            block
            for block in first["blocks"]
            if block["type"] == "text"
            and any(p["text"] == "Text box note" for p in block["paragraphs"])
        )
        self.assertEqual(textbox_block["role"], "body")
        self.assertGreaterEqual(textbox_block["shape_index"], 0)

    def test_blank_slide(self):
        blank = self.slides[2]
        self.assertEqual(blank["title"], "")
        self.assertEqual(blank["blocks"], [])

    def test_picture_is_placeholder_not_invented_content(self):
        picture = next(
            block
            for block in self.slides[1]["blocks"]
            if block["type"] == "unsupported" and block["shape_type"] == "PICTURE"
        )
        self.assertIn("not parsed", picture["note"])
        self.assertNotIn("text", picture)

    def test_special_characters_and_html_escaping(self):
        special = self.slides[3]
        self.assertEqual(
            special["title"], 'Special <script>alert("x")</script> & "quotes"'
        )
        self.assertIn("&lt;script&gt;", self.html_text)
        self.assertIn("&amp;", self.html_text)
        self.assertIn("&lt;td&gt; &amp; &quot;cell&quot;", self.html_text)
        self.assertNotIn("<script>alert", self.html_text)

    def test_deterministic_output(self):
        again = parse_pptx(self.pptx_path)
        self.assertEqual(dumps_content(again), self.json_text)
        self.assertEqual(render_html(again), self.html_text)
        self.assertEqual(again["source_sha256"], self.source_sha)

    def test_output_matches_schema(self):
        schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
        validator = jsonschema.Draft202012Validator(schema)
        errors = sorted(validator.iter_errors(self.content), key=lambda error: list(error.path))
        self.assertEqual(errors, [])

    def test_invalid_pptx_has_clear_error(self):
        with TemporaryDirectory() as folder:
            bad = Path(folder) / "bad.pptx"
            bad.write_text("not a pptx", encoding="utf-8")
            with self.assertRaises(PptxContentError) as raised:
                parse_pptx(bad)
            self.assertIn("INVALID_PPTX", str(raised.exception))

            result = subprocess.run(
                [sys.executable, str(SCRIPT_PATH), str(bad), "--outdir", str(Path(folder) / "out")],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 2)
            self.assertIn("INVALID_PPTX", result.stderr)

    def test_cli_outputs_json_html_and_no_filename_leak(self):
        with TemporaryDirectory() as folder:
            outdir = Path(folder) / "out"
            result = subprocess.run(
                [sys.executable, str(SCRIPT_PATH), str(self.pptx_path), "--outdir", str(outdir)],
                cwd=REPO_ROOT,
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            json_path = outdir / "content.json"
            html_path = outdir / "preview.html"
            self.assertTrue(json_path.is_file())
            self.assertTrue(html_path.is_file())

            json_text = json_path.read_text(encoding="utf-8")
            html_text = html_path.read_text(encoding="utf-8")
            parsed = json.loads(json_text)
            self.assertEqual(parsed["source_sha256"], self.source_sha)
            self.assertNotIn(self.pptx_path.name, json_text)
            self.assertNotIn(str(self.pptx_path), json_text)
            self.assertNotIn(self.pptx_path.name, html_text)
            self.assertNotIn(str(self.pptx_path), html_text)
            self.assertNotIn(self.pptx_path.name, result.stdout)
            self.assertNotIn("FEISHU_USER_ACCESS_TOKEN", json_text)
            self.assertNotIn("<script>alert", html_text)


if __name__ == "__main__":
    unittest.main()
