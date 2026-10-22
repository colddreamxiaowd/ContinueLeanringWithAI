"""Deterministic PPTX content extraction for the V4.2 content pipeline.

The output is a versioned, traceable JSON document plus a standalone local HTML
preview. No content is uploaded or published, and no input filename or path is
stored in the output.
"""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

SCHEMA_VERSION = "pptx-content-v1"
TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
UNSUPPORTED_SHAPE_TYPES = {
    getattr(MSO_SHAPE_TYPE, name)
    for name in (
        "PICTURE",
        "LINKED_PICTURE",
        "CHART",
        "MEDIA",
        "OLE_OBJECT",
        "LINKED_OLE_OBJECT",
        "GROUP",
    )
    if hasattr(MSO_SHAPE_TYPE, name)
}


class PptxContentError(Exception):
    """Raised when a PPTX cannot be read or converted safely."""


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _shape_type_name(shape: Any) -> str:
    shape_type = getattr(shape, "shape_type", None)
    if shape_type is None:
        return "UNKNOWN"
    return getattr(shape_type, "name", None) or str(shape_type)


def _text_block(shape: Any, shape_index: int, role: str) -> dict[str, Any] | None:
    paragraphs = []
    for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs):
        paragraphs.append(
            {
                "index": paragraph_index,
                "level": int(paragraph.level or 0),
                "text": paragraph.text or "",
            }
        )
    if not any(paragraph["text"].strip() for paragraph in paragraphs):
        return None
    return {
        "type": "text",
        "role": role,
        "shape_index": shape_index,
        "paragraphs": paragraphs,
    }


def _table_block(shape: Any, shape_index: int) -> dict[str, Any]:
    return {
        "type": "table",
        "shape_index": shape_index,
        "rows": [[cell.text or "" for cell in row.cells] for row in shape.table.rows],
    }


def _unsupported_block(shape: Any, shape_index: int) -> dict[str, Any] | None:
    shape_type = _shape_type_name(shape)
    if getattr(shape, "has_chart", False):
        shape_type = "CHART"
        note = "Chart data is not parsed in v1."
    elif shape.shape_type in UNSUPPORTED_SHAPE_TYPES:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            note = "Grouped shapes are not traversed in v1."
        elif shape.shape_type in {
            getattr(MSO_SHAPE_TYPE, "PICTURE", None),
            getattr(MSO_SHAPE_TYPE, "LINKED_PICTURE", None),
        }:
            note = "Picture content is not parsed in v1."
        elif shape.shape_type == getattr(MSO_SHAPE_TYPE, "MEDIA", None):
            note = "Embedded media is not parsed in v1."
        else:
            note = "Embedded object is not parsed in v1."
    else:
        return None
    return {
        "type": "unsupported",
        "shape_index": shape_index,
        "shape_type": shape_type,
        "note": note,
    }


def _slide_document(slide: Any, slide_index: int) -> dict[str, Any]:
    title_shape = slide.shapes.title
    title_shape_id = getattr(title_shape, "shape_id", None)
    blocks: list[dict[str, Any]] = []
    title_text = ""

    for shape_index, shape in enumerate(slide.shapes):
        try:
            if shape.has_table:
                blocks.append(_table_block(shape, shape_index))
                continue

            if shape.has_text_frame:
                is_title = title_shape_id is not None and shape.shape_id == title_shape_id
                role = "title" if is_title else "body"
                block = _text_block(shape, shape_index, role)
                if block is not None:
                    blocks.append(block)
                    if is_title and not title_text:
                        title_text = "\n".join(
                            paragraph["text"] for paragraph in block["paragraphs"]
                        ).strip()
                continue

            unsupported = _unsupported_block(shape, shape_index)
            if unsupported is not None:
                blocks.append(unsupported)
        except Exception as exc:  # keep one malformed shape from aborting the deck
            blocks.append(
                {
                    "type": "unsupported",
                    "shape_index": shape_index,
                    "shape_type": _shape_type_name(shape),
                    "note": f"Shape could not be read ({type(exc).__name__}).",
                }
            )

    return {
        "slide_index": slide_index,
        "slide_number": slide_index + 1,
        "title": title_text,
        "blocks": blocks,
    }


def parse_pptx(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    if not source.is_file():
        raise PptxContentError("INVALID_PPTX: input file does not exist")

    digest = sha256_file(source)
    try:
        presentation = Presentation(str(source))
    except Exception as exc:  # python-pptx raises several exception types
        raise PptxContentError(
            f"INVALID_PPTX: cannot open presentation ({type(exc).__name__})"
        ) from None

    try:
        slides = [
            _slide_document(slide, slide_index)
            for slide_index, slide in enumerate(presentation.slides)
        ]
    except Exception as exc:
        raise PptxContentError(
            f"INVALID_PPTX: cannot read slide content ({type(exc).__name__})"
        ) from None

    return {
        "schema_version": SCHEMA_VERSION,
        "source_sha256": digest,
        "slides": slides,
    }


def dumps_content(content: dict[str, Any]) -> str:
    """Return a stable JSON representation suitable for committing or diffing."""
    return json.dumps(content, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def render_html(content: dict[str, Any]) -> str:
    """Render a standalone local preview. All source text is HTML-escaped."""
    digest = html.escape(content["source_sha256"])
    parts = [
        "<!doctype html>",
        '<html lang="zh-CN">',
        "<head>",
        '<meta charset="utf-8">',
        "<title>PPTX Content Preview</title>",
        "<style>",
        "body{font-family:system-ui,-apple-system,'Segoe UI',sans-serif;max-width:960px;"
        "margin:2rem auto;padding:0 1rem;line-height:1.65;color:#1f2328}",
        ".slide{border:1px solid #d0d7de;border-radius:8px;padding:1rem 1.5rem;"
        "margin:1.5rem 0}",
        ".block{margin:1rem 0}",
        ".text p{margin:.25rem 0;white-space:pre-wrap}",
        ".level-1{margin-left:1.5rem}.level-2{margin-left:3rem}"
        ".level-3{margin-left:4.5rem}.level-4{margin-left:6rem}"
        ".level-5{margin-left:7.5rem}.level-6{margin-left:9rem}"
        ".level-7{margin-left:10.5rem}.level-8{margin-left:12rem}",
        "table{border-collapse:collapse;width:100%}",
        "td{border:1px solid #d0d7de;padding:.5rem;vertical-align:top;"
        "white-space:pre-wrap}",
        ".unsupported{color:#8a5a00;background:#fff8e1;padding:.5rem;"
        "border-radius:4px}",
        ".empty{color:#57606a;font-style:italic}",
        "footer{margin-top:2rem;color:#57606a;font-size:.85rem}",
        "</style>",
        "</head>",
        "<body>",
        "<h1>PPTX Content Preview</h1>",
    ]

    for slide in content["slides"]:
        title = html.escape(slide["title"])
        parts.append(f'<section class="slide" id="slide-{slide["slide_number"]}">')
        heading = f'Slide {slide["slide_number"]}'
        if title:
            heading += f": {title}"
        parts.append(f"<h2>{heading}</h2>")

        if not slide["blocks"]:
            parts.append('<p class="empty">Blank slide.</p>')

        for block in slide["blocks"]:
            if block["type"] == "text":
                # Title text is already rendered in the slide heading. Keep the
                # title block in JSON for traceability, but avoid duplicate UI.
                if block["role"] == "title" and slide["title"]:
                    continue
                parts.append(f'<div class="block text {block["role"]}">')
                for paragraph in block["paragraphs"]:
                    escaped = html.escape(paragraph["text"])
                    if not escaped:
                        escaped = "&nbsp;"
                    parts.append(
                        f'<p class="level-{paragraph["level"]}">{escaped}</p>'
                    )
                parts.append("</div>")
            elif block["type"] == "table":
                parts.append('<div class="block table"><table>')
                for row in block["rows"]:
                    cells = "".join(
                        f"<td>{html.escape(cell) if cell else '&nbsp;'}</td>"
                        for cell in row
                    )
                    parts.append(f"<tr>{cells}</tr>")
                parts.append("</table></div>")
            elif block["type"] == "unsupported":
                parts.append(
                    '<p class="unsupported">Unsupported shape '
                    f'{block["shape_index"]}: {html.escape(block["shape_type"])}'
                    f' — {html.escape(block["note"])}</p>'
                )
        parts.append("</section>")

    parts.extend(
        [
            f"<footer>Source SHA-256: <code>{digest}</code><br>"
            "Generated locally; not for publication.</footer>",
            "</body>",
            "</html>",
        ]
    )
    return "\n".join(parts) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Convert a local PPTX into deterministic JSON and a local HTML preview."
    )
    parser.add_argument("pptx", type=Path, help="local .pptx file")
    parser.add_argument(
        "--outdir",
        type=Path,
        default=Path("build/v42"),
        help="directory for content.json and preview.html (default: build/v42)",
    )
    args = parser.parse_args(argv)

    try:
        content = parse_pptx(args.pptx)
        json_text = dumps_content(content)
        html_text = render_html(content)
        args.outdir.mkdir(parents=True, exist_ok=True)
        json_path = args.outdir / "content.json"
        html_path = args.outdir / "preview.html"
        json_path.write_text(json_text, encoding="utf-8")
        html_path.write_text(html_text, encoding="utf-8")
    except PptxContentError as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        return 2
    except OSError:
        print("ERROR OUTPUT_IO: cannot write output files", file=sys.stderr)
        return 2

    block_count = sum(len(slide["blocks"]) for slide in content["slides"])
    print(
        f"OK slides={len(content['slides'])} blocks={block_count} "
        f"json={json_path} html={html_path}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
